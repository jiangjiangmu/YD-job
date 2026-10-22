# -*- coding: utf-8 -*-
"""YD-job 岗位匹配分析脚本

读取 data/raw_jobs.json（spider.py 采集结果），
按 config.yaml 中 profile 画像对每个岗位加权打分，
导出 Excel 报告到 data/ 目录，并在控制台输出 Top 摘要。

打分模型（满分 100，另可扣分）：
    技能匹配 40 + 薪资匹配 25 + 经验匹配 15 + 学历匹配 10 + 区域偏好 10
    岗位名命中排除词每项 -20
兼职/实习默认不参与打分（profile.exclude_part_time 控制），单独归入附表。

用法：
    .venv\\Scripts\\python.exe analyzer.py
"""
import json
import os
import re
import sys
from collections import Counter
from datetime import datetime

import yaml
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # scripts/ -> project root
CONFIG_PATH = os.path.join(BASE_DIR, "config.yaml")
JOBS_PATH = os.path.join(BASE_DIR, "data", "raw_jobs.json")

WEIGHTS = {"skill": 40, "salary": 25, "experience": 15, "degree": 10, "district": 10}
SKILL_ALIAS = {"py": "python", "js": "javascript", "ts": "typescript", "golang": "go"}
DEGREE_LEVELS = [("博士", 5), ("硕士", 4), ("本科", 3), ("大专", 2), ("中专", 1), ("中技", 1), ("高中", 1), ("学历不限", 0), ("初中", 0)]

DEFAULT_PROFILE = {
    "desired_skills": [],
    "years_of_experience": None,
    "min_salary_k": None,
    "max_salary_k": None,
    "own_degree": "大专",
    "preferred_districts": [],
    "excluded_keywords": [],
    "exclude_part_time": True,
    "notes": "",
}


def parse_expected_salary(text):
    """解析 '8K-12K' / '8-12K' / '12K以上' 形式的期望薪资，返回 (最低K, 最高K)"""
    s = str(text or "").strip()
    m = re.search(r"([\d.]+)\s*[Kk]?\s*[-~]\s*([\d.]+)\s*[Kk]?", s)
    if m:
        return float(m.group(1)), float(m.group(2))
    m = re.search(r"([\d.]+)\s*[Kk]\s*以上", s)
    if m:
        return float(m.group(1)), None
    return None, None


def load_profile():
    with open(CONFIG_PATH, "r", encoding="utf-8") as f:
        cfg = yaml.safe_load(f) or {}
    raw = dict(cfg.get("profile") or {})
    p = dict(DEFAULT_PROFILE)
    p.update({k: v for k, v in raw.items() if k in p})
    # 兼容 job-desktop 指令的字段名：skills / experience_years / expected_salary / preferred_regions / notes
    if raw.get("skills"):
        p["desired_skills"] = raw["skills"]
    if raw.get("experience_years") is not None:
        p["years_of_experience"] = raw["experience_years"]
    if raw.get("expected_salary"):
        p["min_salary_k"], p["max_salary_k"] = parse_expected_salary(raw["expected_salary"])
    if raw.get("preferred_regions"):
        p["preferred_districts"] = raw["preferred_regions"]
    if raw.get("notes") is not None:
        p["notes"] = raw["notes"]
    return p


def load_jobs():
    if not os.path.exists(JOBS_PATH):
        print("[错误] 找不到 data/raw_jobs.json，请先运行: .venv\\Scripts\\python.exe spider.py")
        return None
    with open(JOBS_PATH, "r", encoding="utf-8") as f:
        data = json.load(f)
    if isinstance(data, dict):
        return data.get("jobs") or [], data.get("meta") or {}
    return data, {}


# ---------- 薪资 ----------

def parse_salary(desc):
    """返回 (最低K, 最高K, 类别)；类别: k / daily / negotiable / unknown"""
    s = str(desc or "").strip()
    if not s or "面议" in s:
        return None, None, "negotiable"
    if re.search(r"元/(天|周|小时)", s):
        return None, None, "daily"
    m = re.search(r"([\d.]+)\s*[Kk]?\s*[-~]\s*([\d.]+)\s*[Kk]", s)
    if m:
        return float(m.group(1)), float(m.group(2)), "k"
    m = re.search(r"([\d.]+)\s*万\s*[-~]\s*([\d.]+)\s*万", s)
    if m:
        return float(m.group(1)) * 10, float(m.group(2)) * 10, "k"
    m = re.search(r"([\d.]+)\s*千\s*[-~]\s*([\d.]+)\s*万", s)
    if m:
        return float(m.group(1)), float(m.group(2)) * 10, "k"
    m = re.search(r"([\d.]+)\s*[-~]\s*([\d.]+)\s*万", s)
    if m:
        return float(m.group(1)) * 10, float(m.group(2)) * 10, "k"
    m = re.search(r"([\d.]+)\s*[Kk]\s*[以+]", s)
    if m:
        return float(m.group(1)), None, "k"
    return None, None, "unknown"


def salary_score(lo, hi, kind, exp_min, exp_max):
    w = WEIGHTS["salary"]
    if kind in ("negotiable", "unknown"):
        return w * 0.5
    if lo is None:
        return w * 0.5
    if hi is None:
        hi = lo * 1.3
    if exp_min is None or exp_min <= 0:
        return w
    if lo >= exp_min:
        return w
    if hi <= exp_min:
        ratio = hi / exp_min
        return max(0.0, min(0.5, (ratio - 0.5) / 0.5)) * w
    overlap = hi - exp_min
    span = max(1e-6, hi - lo)
    frac = min(1.0, overlap / span)
    return (0.5 + 0.5 * frac) * w


# ---------- 经验 ----------

def parse_experience(text):
    t = str(text or "").strip()
    if not t or "不限" in t:
        return None, None
    if "应届" in t or "在校" in t:
        return 0.0, 1.0
    m = re.search(r"(\d+)\s*[-~]\s*(\d+)\s*年", t)
    if m:
        return float(m.group(1)), float(m.group(2))
    m = re.search(r"(\d+)\s*年以上", t)
    if m:
        return float(m.group(1)), None
    return None, None


def experience_score(job_exp, own_years):
    w = WEIGHTS["experience"]
    lo, hi = parse_experience(job_exp)
    if lo is None:
        return w * 0.8, "经验不限"
    if own_years is None:
        return w * 0.5, ""
    if own_years >= lo and (hi is None or own_years <= hi):
        return w, ""
    if own_years < lo:
        return max(1.0, own_years / lo * w), f"要求{job_exp}"
    return w * max(0.6, (hi or 0) / own_years), "经验超要求"


# ---------- 学历 ----------

def degree_level(text):
    t = str(text or "")
    for name, lv in DEGREE_LEVELS:
        if name in t:
            return lv
    return -1


def degree_score(job_degree, own_degree):
    w = WEIGHTS["degree"]
    req = degree_level(job_degree)
    own = degree_level(own_degree)
    if own <= 0:
        return w * 0.6, ""
    if req < 0:
        return w * 0.6, ""
    if req <= own:
        return w, ""
    partial = w * own / req if req > 0 else 0.0
    return partial, f"要求{job_degree}高于本人{own_degree}"


# ---------- 区域 ----------

def district_score(job, prefs):
    w = WEIGHTS["district"]
    d = str(job.get("district") or "").strip()
    b = str(job.get("business_district") or "").strip()
    if not prefs:
        return w * 0.6, ""
    for p in prefs:
        if p in d or p in b:
            return w, ""
    if d:
        return w * 0.3, f"{d or b}"
    return w * 0.5, ""


# ---------- 技能 ----------

def norm_skill(x):
    s = str(x).strip().lower()
    return SKILL_ALIAS.get(s, s)


def match_skills(desired, job):
    corpus = [norm_skill(x) for x in (job.get("skills") or [])]
    name = norm_skill(job.get("job_name") or "")
    matched, missing = [], []
    for sk in desired or []:
        q = norm_skill(sk)
        hit = any(q and (q in c) for c in corpus) or (q and q in name)
        (matched if hit else missing).append(sk)
    return matched, missing


# ---------- 兼职/实习 ----------

def is_part_time(job):
    sal = str(job.get("salary_desc") or "")
    name = str(job.get("job_name") or "")
    labels = [str(x) for x in (job.get("job_labels") or [])]
    if re.search(r"元/(天|周|小时)", sal):
        return True
    if re.search(r"实习|兼职|暑假|寒假|在校生|大学生", name):
        return True
    for lb in labels:
        if re.search(r"\d+\s*天/周", lb) or re.search(r"^\d+个月$", lb):
            return True
    return False


# ---------- 综合打分 ----------

def score_job(job, profile):
    lo, hi, kind = parse_salary(job.get("salary_desc"))
    matched, missing = match_skills(profile.get("desired_skills"), job)
    n_desired = len(profile.get("desired_skills") or [])
    s_skill = WEIGHTS["skill"] * (len(matched) / n_desired) if n_desired else WEIGHTS["skill"] * 0.5
    s_salary = salary_score(lo, hi, kind, profile.get("min_salary_k"), profile.get("max_salary_k"))
    s_exp, note_exp = experience_score(job.get("experience"), profile.get("years_of_experience"))
    s_deg, note_deg = degree_score(job.get("degree"), profile.get("own_degree"))
    s_dis, note_dis = district_score(job, profile.get("preferred_districts"))

    notes = [n for n in (note_exp, note_deg) if n]
    penalty = 0.0
    for kw in profile.get("excluded_keywords") or []:
        if kw and kw in str(job.get("job_name") or ""):
            penalty -= 20
            notes.append(f"排除词:{kw}")
    if kind in ("negotiable", "unknown"):
        notes.append("薪资未标注")

    total = max(0.0, min(100.0, s_skill + s_salary + s_exp + s_deg + s_dis + penalty))
    return {
        "job": job,
        "total": round(total, 1),
        "s_skill": round(s_skill, 1),
        "s_salary": round(s_salary, 1),
        "s_exp": round(s_exp, 1),
        "s_deg": round(s_deg, 1),
        "s_dis": round(s_dis, 1),
        "penalty": penalty,
        "matched": matched,
        "missing": missing,
        "sal_lo": lo,
        "sal_hi": hi,
        "sal_kind": kind,
        "notes": notes,
    }


# ---------- 投递建议 ----------

BANDS = ["⭐必投", "去", "可去", "缓", "不去"]


def build_advice(r, profile):
    """根据打分明细生成单岗位文字分析与投递建议档位，返回 (分析文本, 档位)"""
    j = r["job"]
    pros, cons = [], []
    if r["matched"]:
        pros.append("对口:" + "/".join(r["matched"][:4]))
    else:
        cons.append("技能标签未直接命中")
    if r["sal_kind"] == "k" and r["sal_lo"] is not None:
        pros.append(f"薪资{r['sal_lo']:.0f}-{r['sal_hi']:.0f}K达标")
    else:
        cons.append("薪资面议/未标注")
    if r["s_exp"] >= WEIGHTS["experience"]:
        pros.append("经验匹配")
    else:
        cons.append(f"经验要求[{j.get('experience') or '不限'}]")
    own_lv = degree_level(profile.get("own_degree"))
    req_lv = degree_level(j.get("degree"))
    degree_blocked = own_lv > 0 and req_lv > own_lv
    if degree_blocked:
        cons.append(f"学历门槛[{j.get('degree')}]")
    else:
        pros.append(f"学历OK[{j.get('degree') or '不限'}]")
    d = j.get("district") or ""
    prefs = profile.get("preferred_districts") or []
    if d:
        if d in prefs:
            pros.append(f"{d}通勤友好")
        elif prefs:
            cons.append(f"{d}不在偏好区")
    scale = j.get("company_scale") or ""
    if any(s in scale for s in ("1000-9999", "10000")):
        pros.append("规模企业")
    text = ("✓ " + "；".join(pros)) if pros else ""
    if cons:
        text += ("  ⚠ " if text else "⚠ ") + "；".join(cons)
    if r["total"] >= 70:
        b = 0
    elif r["total"] >= 65:
        b = 1
    elif r["total"] >= 58:
        b = 2
    elif r["total"] >= 48:
        b = 3
    else:
        b = 4
    if degree_blocked and b < 4:
        b += 1
    return text or "-", BANDS[b]


# ---------- Excel ----------

HEADER = ["排名", "综合得分", "要不要去", "匹配分析(为什么)", "技能(40)", "薪资(25)", "经验(15)", "学历(10)", "区域(10)",
          "命中技能", "待补技能", "职位名称", "薪资", "月薪区间(K)", "公司", "行业", "公司规模",
          "行政区", "商圈", "经验要求", "学历要求", "HR", "HR头衔", "岗位标签", "备注", "岗位链接"]
WIDTHS = [5, 8, 9, 46, 8, 8, 8, 8, 8, 20, 18, 30, 12, 11, 18, 13, 10, 9, 9, 9, 9, 9, 9, 15, 13, 42]

HEADER_FILL = PatternFill("solid", fgColor="2F5597")
HEADER_FONT = Font(bold=True, color="FFFFFF")
LINK_FONT = Font(color="0563C1", underline="single")
WRAP_COLS = {4, 10, 11, 24}  # 1-based 列号：匹配分析/命中技能/待补技能/岗位标签

BAND_COLORS = {
    "⭐必投": "C6EFCE",
    "去": "E2EFDA",
    "可去": "FFEB9C",
    "缓": "FCD5B4",
    "不去": "FFC7CE",
}


def write_scored_sheet(ws, rows, profile):
    ws.append(HEADER)
    for c in range(1, len(HEADER) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[cell.column_letter].width = WIDTHS[c - 1]
    for i, r in enumerate(rows, start=1):
        j = r["job"]
        sal_range = f"{r['sal_lo']:.0f}-{r['sal_hi']:.0f}" if r["sal_lo"] is not None and r["sal_kind"] == "k" else "见薪资列"
        link = j.get("job_url") or ""
        advice_text, band = build_advice(r, profile)
        data = [i, r["total"], band, advice_text,
                r["s_skill"], r["s_salary"], r["s_exp"], r["s_deg"], r["s_dis"],
                ", ".join(r["matched"]), ", ".join(r["missing"]),
                j.get("job_name", ""), j.get("salary_desc", ""), sal_range,
                j.get("company", ""), j.get("industry", ""), j.get("company_scale", ""),
                j.get("district", ""), j.get("business_district", ""),
                j.get("experience", ""), j.get("degree", ""),
                j.get("boss_name", ""), j.get("boss_title", ""),
                ", ".join(j.get("job_labels") or []), "; ".join(r["notes"]), link]
        ws.append(data)
        row_idx = ws.max_row
        for c in WRAP_COLS:
            ws.cell(row=row_idx, column=c).alignment = Alignment(wrap_text=True, vertical="top")
        band_cell = ws.cell(row=row_idx, column=3)
        color = BAND_COLORS.get(band)
        if color:
            band_cell.fill = PatternFill("solid", fgColor=color)
        link_cell = ws.cell(row=row_idx, column=len(HEADER))
        if link:
            link_cell.hyperlink = link
            link_cell.font = LINK_FONT
    ws.freeze_panes = "D2"
    if rows:
        ws.conditional_formatting.add(
            f"B2:B{ws.max_row}",
            ColorScaleRule(start_type="min", start_color="F8696B",
                          mid_type="percentile", mid_value=50, mid_color="FFEB84",
                          end_type="max", end_color="63BE7B"))


def write_parttime_sheet(ws, jobs):
    header = ["职位名称", "薪资", "公司", "行业", "公司规模", "行政区", "商圈", "岗位标签", "岗位链接"]
    widths = [30, 12, 20, 14, 10, 10, 10, 18, 42]
    ws.append(header)
    for c in range(1, len(header) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        ws.column_dimensions[cell.column_letter].width = widths[c - 1]
    for j in jobs:
        link = j.get("job_url") or ""
        ws.append([j.get("job_name", ""), j.get("salary_desc", ""), j.get("company", ""),
                   j.get("industry", ""), j.get("company_scale", ""), j.get("district", ""),
                   j.get("business_district", ""), ", ".join(j.get("job_labels") or []), link])
        lc = ws.cell(row=ws.max_row, column=len(header))
        if link:
            lc.hyperlink = link
            lc.font = LINK_FONT
    ws.freeze_panes = "A2"


def top_comment(r, profile):
    """为 Top 岗位生成规则化点评文字"""
    j = r["job"]
    parts = []
    if r["matched"]:
        parts.append("对口技能：" + "、".join(r["matched"][:4]))
    sal = str(j.get("salary_desc") or "")
    if r["sal_kind"] == "k" and r["sal_lo"] is not None:
        lo = r["sal_lo"]
        hi = r["sal_hi"] or r["sal_lo"]
        tag = "符合期望区间" if lo >= 8 else "略低于期望"
        extra = ""
        m = re.search(r"[·•](\d+)薪", sal)
        if m:
            n = int(m.group(1))
            extra = f"，年包约 {lo * n / 10:.1f}-{hi * n / 10:.1f} 万"
        parts.append(f"月薪 {lo:.0f}-{hi:.0f}K{extra}，{tag}")
    scale = j.get("company_scale") or ""
    if "10000" in scale or "1000-9999" in scale:
        parts.append("规模企业，平台稳定")
    name = j.get("job_name", "")
    if "AI" in name or "人工智能" in name:
        parts.append("AI 方向，契合转型路线")
    exp = j.get("experience") or ""
    if not exp or "不限" in exp:
        parts.append("经验不限，你的年限富余")
    own_lv = degree_level(profile.get("own_degree"))
    if degree_level(j.get("degree")) <= own_lv:
        parts.append(f"学历要求「{j.get('degree') or '不限'}」可投")
    else:
        parts.append(f"学历要求「{j.get('degree')}」有门槛，投递前权衡")
    d = j.get("district") or ""
    if d:
        prefs = profile.get("preferred_districts") or []
        parts.append(f"位于{d}" + ("，通勤友好" if d in prefs else "，不在偏好区"))
    return "；".join(parts) + "。"


def write_conclusion_sheet(ws, meta, profile, rows, part_jobs):
    """分析结论页：打分模型 / 总体结论 / 投递策略 / Top 点评 / 简历对照 / 使用说明"""
    ws.column_dimensions["A"].width = 100

    def para(text, bold=False, prefix=""):
        for seg in str(text).split("\n"):
            seg = prefix + seg
            while seg:
                ws.append([seg[:48]])
                c = ws.cell(row=ws.max_row, column=1)
                c.alignment = Alignment(vertical="top")
                if bold:
                    c.font = Font(bold=True)
                seg = seg[48:]

    n = len(rows)
    if not n:
        para("无打分数据，请先运行 spider.py 采集。")
        return

    bands = Counter(build_advice(r, profile)[1] for r in rows)
    own_lv = degree_level(profile.get("own_degree"))
    prefs = set(profile.get("preferred_districts") or [])
    sal_ok = deg_sal = full = 0
    for r in rows:
        ok_sal = r["sal_kind"] == "k" and r["sal_lo"] is not None and (r["sal_lo"] + (r["sal_hi"] or r["sal_lo"])) / 2 >= 8
        ok_deg = degree_level(r["job"].get("degree")) <= own_lv
        ok_dis = (r["job"].get("district") or "") in prefs
        if ok_sal:
            sal_ok += 1
        if ok_sal and ok_deg:
            deg_sal += 1
        if ok_sal and ok_deg and ok_dis:
            full += 1
    deg_c = Counter((r["job"].get("degree") or "不限").strip() or "不限" for r in rows)
    exp_c = Counter((r["job"].get("experience") or "不限").strip() or "不限" for r in rows)
    dis_c = Counter(r["job"].get("district") or "未知" for r in rows)

    para("一、打分模型（满分 100）", bold=True)
    para("技能匹配 40 + 薪资匹配 25 + 经验匹配 15 + 学历匹配 10 + 区域偏好 10；岗位名命中排除词每项 -20；兼职/实习不参与打分，单独归表。")
    para("读分方法：对口岗位的薪资/经验/学历/区域四项通常全满（60 分做底），技能项因直聘标签多为描述性短语（如「系统运维」）普遍得 5-13 分。因此 65 分以上 = 硬条件全部达标；58-65 分多为区域或学历单点扣分，不代表岗位差。")
    para("")

    para("二、总体结论", bold=True)
    para(f"本次参与打分 {n} 条（另有兼职/实习 {len(part_jobs)} 条单独归表）。")
    para(f"分档统计：⭐必投 {bands.get('⭐必投', 0)} 条 | 去 {bands.get('去', 0)} 条 | 可去 {bands.get('可去', 0)} 条 | 缓 {bands.get('缓', 0)} 条 | 不去 {bands.get('不去', 0)} 条。")
    para(f"可投矩阵：薪资达标 {sal_ok} 条 → 加上学历可投 {deg_sal} 条 → 再加偏好区域 {full} 条（三满足，可直接投）。")
    ok_deg_n = deg_c.get("大专", 0) + deg_c.get("学历不限", 0) + deg_c.get("中专/中技", 0) + deg_c.get("高中", 0)
    para(f"学历：大专可投 {deg_c.get('大专', 0)} 条 + 不限 {deg_c.get('学历不限', 0)} 条等，合计 {ok_deg_n} 条（约 {ok_deg_n * 100 // n}%），本赛道对大专友好。")
    para(f"经验：3-5 年要求 {exp_c.get('3-5年', 0)} 条、1-3 年 {exp_c.get('1-3年', 0)} 条，你的 {profile.get('years_of_experience')} 年经验正覆盖主力区间。")
    para("区域分布 Top4：" + "、".join(f"{k}{v}条" for k, v in dis_c.most_common(4)) + "。")
    para("")

    para("三、投递策略", bold=True)
    para(f"第一批（今天就投）：⭐必投 {bands.get('⭐必投', 0)} 条——见「岗位总表」绿色行，硬条件全达标。")
    para(f"第二批（本周铺量）：「去」+「可去」共 {bands.get('去', 0) + bands.get('可去', 0)} 条，优先宝安/南山/龙华的大专可投岗，每天 8-10 条分批沟通（直聘活跃度影响曝光权重，不要一次性全发）。")
    para(f"暂缓：「缓」{bands.get('缓', 0)} 条；跳过：「不去」{bands.get('不去', 0)} 条及本科硬门槛且技能不对口的岗位（「备注」列有扣分原因）。")
    para("")

    para("四、Top 10 岗位点评", bold=True)
    for i, r in enumerate(rows[:10], 1):
        j = r["job"]
        para(f"#{i} [{r['total']}分 · {build_advice(r, profile)[1]}] {j.get('job_name', '')} | {j.get('company', '')} | {j.get('salary_desc', '')} | {j.get('district', '')}", bold=True)
        para(top_comment(r, profile), prefix="    ")
    para("")

    para("五、简历关键词对照（岗位高频要求 vs 你的技能）", bold=True)
    counter = Counter()
    for r in rows:
        for s in r["job"].get("skills") or []:
            counter[str(s).strip()] += 1
    own = [str(x) for x in (profile.get("desired_skills") or [])]
    have, gap = [], []
    for name, cnt in counter.most_common(15):
        hit = any(o.lower() in name.lower() or name.lower() in o.lower() for o in own)
        (have if hit else gap).append(f"{name}({cnt})")
    para("你已具备（确保简历原样出现这些词，HR 搜索友好）：" + ("、".join(have) if have else "无"))
    para("岗位常要求但你画像未覆盖（酌情补强或面试前预习）：" + ("、".join(gap) if gap else "无"))
    para("差异化提示：Python/自动化/Docker/AI 相关要求在本数据集出现频次不低，你的自动化项目经历放在简历显眼位置即是加分项。")
    para("")

    para("六、使用说明", bold=True)
    para("1. 「岗位总表」按综合得分降序，26 列全量：C 列「要不要去」（带底色）、D 列匹配分析（为什么）、E-I 列五维分项评分、J-K 列命中/待补技能，职位信息与链接齐全；最后一列点击直达直聘岗位页。")
    para("2. 「兼职实习(未打分)」为被过滤的实习/兼职岗位。")
    para("3. 「统计与配置」含数据集统计、高频技能需求与当前打分画像回显。")
    para("4. 更换关键词/城市/画像：编辑 config.yaml 后依次运行 spider.py、analyzer.py。")


def write_stats_sheet(ws, meta, profile, rows, part_jobs):
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 46
    items = [
        ("== 采集信息 ==", ""),
        ("关键词 / 城市编码", f"{'、'.join(map(str, meta.get('keywords') or [meta.get('keyword') or '']))} / {meta.get('city_code', '')}"),
        ("采集时间", str(meta.get("fetched_at", ""))),
        ("采集总数 / 去重后", f"{meta.get('total_collected', '')} / {meta.get('total_unique', '')}"),
        ("", ""),
        ("== 分析结果 ==", ""),
        ("参与打分岗位数", len(rows)),
        ("兼职/实习（未打分）", len(part_jobs)),
        ("平均分", f"{sum(r['total'] for r in rows) / len(rows):.1f}" if rows else "-"),
        ("≥70 分岗位数", sum(1 for r in rows if r['total'] >= 70)),
        ("50-70 分岗位数", sum(1 for r in rows if 50 <= r['total'] < 70)),
        ("", ""),
        ("== 高频技能需求（Top15）==", ""),
    ]
    counter = Counter()
    for j in (r["job"] for r in rows):
        for s in j.get("skills") or []:
            counter[str(s).strip()] += 1
    for name, cnt in counter.most_common(15):
        items.append((name, f"{cnt} 次"))
    items += [("", ""), ("== 数据集画像 ==", "")]
    buckets = {}
    for r in rows:
        lo, kind = r["sal_lo"], r["sal_kind"]
        if kind != "k" or lo is None:
            buckets["未标注/面议"] = buckets.get("未标注/面议", 0) + 1
            continue
        mid = (lo + (r["sal_hi"] or lo)) / 2
        key = "<8K" if mid < 8 else "8-12K" if mid < 12 else "12-15K" if mid < 15 else "15-20K" if mid < 20 else "20K+"
        buckets[key] = buckets.get(key, 0) + 1
    for k in ["<8K", "8-12K", "12-15K", "15-20K", "20K+", "未标注/面议"]:
        if k in buckets:
            items.append((f"薪资 {k}", f"{buckets[k]} 条"))
    deg_c = {}
    for r in rows:
        key = (r["job"].get("degree") or "不限").strip() or "不限"
        deg_c[key] = deg_c.get(key, 0) + 1
    items.append(("学历要求分布", "；".join(f"{k}:{v}" for k, v in sorted(deg_c.items(), key=lambda x: -x[1]))))
    exp_c = {}
    for r in rows:
        key = (r["job"].get("experience") or "不限").strip() or "不限"
        exp_c[key] = exp_c.get(key, 0) + 1
    items.append(("经验要求分布", "；".join(f"{k}:{v}" for k, v in sorted(exp_c.items(), key=lambda x: -x[1]))))
    dis_c = {}
    for r in rows:
        key = r["job"].get("district") or "未知"
        dis_c[key] = dis_c.get(key, 0) + 1
    items.append(("区域分布", "；".join(f"{k}:{v}" for k, v in sorted(dis_c.items(), key=lambda x: -x[1]))))
    own_lv = degree_level(profile.get("own_degree"))
    prefs = set(profile.get("preferred_districts") or [])
    sal_ok = deg_sal = full = 0
    for r in rows:
        ok_sal = r["sal_kind"] == "k" and r["sal_lo"] is not None and (r["sal_lo"] + (r["sal_hi"] or r["sal_lo"])) / 2 >= 8
        ok_deg = degree_level(r["job"].get("degree")) <= own_lv
        ok_dis = (r["job"].get("district") or "") in prefs
        if ok_sal:
            sal_ok += 1
        if ok_sal and ok_deg:
            deg_sal += 1
        if ok_sal and ok_deg and ok_dis:
            full += 1
    items.append(("可投矩阵", f"薪资达标 {sal_ok} | +学历OK {deg_sal} | +区域偏好 {full}"))
    b0 = sum(1 for r in rows if r["total"] >= 70)
    b1 = sum(1 for r in rows if 65 <= r["total"] < 70)
    b2 = sum(1 for r in rows if 58 <= r["total"] < 65)
    b3 = sum(1 for r in rows if 48 <= r["total"] < 58)
    b4 = sum(1 for r in rows if r["total"] < 48)
    items.append(("投递建议分档", f"必投 {b0} | 去 {b1} | 可去 {b2} | 缓 {b3} | 不去 {b4}"))
    items += [("", ""), ("== 个人画像（config.yaml）==", "")]
    for k, v in profile.items():
        items.append((k, ", ".join(map(str, v)) if isinstance(v, list) else str(v)))
    for k, v in items:
        ws.append([k, v])
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=1):
        c = row[0]
        if str(c.value or "").startswith("=="):
            c.font = Font(bold=True)


def main():
    profile = load_profile()
    jobs, meta = load_jobs()
    if jobs is None:
        return 1
    if not jobs:
        print("[错误] raw_jobs.json 里没有岗位数据")
        return 1
    print(f"载入 {len(jobs)} 条岗位 | 画像: 技能={profile.get('desired_skills')} | 年限={profile.get('years_of_experience')} | 期望薪资={profile.get('min_salary_k')}-{profile.get('max_salary_k')}K")

    part_jobs = [j for j in jobs if is_part_time(j)] if profile.get("exclude_part_time", True) else []
    score_jobs_pool = [j for j in jobs if j not in part_jobs]
    rows = sorted((score_job(j, profile) for j in score_jobs_pool), key=lambda r: -r["total"])

    # 控制台摘要
    print("\n=== Top 15 推荐 ===")
    for i, r in enumerate(rows[:15], start=1):
        j = r["job"]
        print(f"{i:>2}. [{r['total']:>5}] {j.get('job_name','')[:24]} | {j.get('salary_desc','')} | {j.get('company','')[:12]} | {j.get('district','')}")

    # Excel
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "分析结论"
    write_conclusion_sheet(ws0, meta, profile, rows, part_jobs)
    ws1 = wb.create_sheet("岗位总表")
    write_scored_sheet(ws1, rows, profile)
    ws2 = wb.create_sheet("兼职实习(未打分)")
    write_parttime_sheet(ws2, part_jobs)
    ws3 = wb.create_sheet("统计与配置")
    write_stats_sheet(ws3, meta, profile, rows, part_jobs)

    out = os.path.join(BASE_DIR, "data", "岗位分析报告.xlsx")
    wb.save(out)

    print(f"\n=== 报告已生成 ===")
    print(f"Excel: {out}")
    print(f"岗位总表 {len(rows)} 条（按得分降序）| 兼职实习 {len(part_jobs)} 条单独归表")
    return 0


if __name__ == "__main__":
    sys.exit(main())


