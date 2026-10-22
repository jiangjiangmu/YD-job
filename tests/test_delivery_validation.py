import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REPORT = os.path.join(ROOT, "data", "岗位分析报告.xlsx")

HEADER_26 = ["排名", "综合得分", "要不要去", "匹配分析(为什么)", "技能(40)", "薪资(25)", "经验(15)", "学历(10)", "区域(10)",
             "命中技能", "待补技能", "职位名称", "薪资", "月薪区间(K)", "公司", "行业", "公司规模",
             "行政区", "商圈", "经验要求", "学历要求", "HR", "HR头衔", "岗位标签", "备注", "岗位链接"]

BAND_COLORS = {"C6EFCE", "E2EFDA", "FFEB9C", "FCD5B4", "FFC7CE"}

def check_fail(msg):
    raise AssertionError(msg)

def test_report_exists():
    assert os.path.exists(REPORT), f"报告缺失: {REPORT}"
    assert os.path.getsize(REPORT) > 0, "报告为空文件"

def test_sheet_structure():
    import openpyxl
    wb = openpyxl.load_workbook(REPORT, read_only=True)
    expected = ["分析结论", "岗位总表", "兼职实习(未打分)", "统计与配置"]
    assert wb.sheetnames == expected, f"Sheet 名/顺序不符: {wb.sheetnames}"

def test_total_sheet_header():
    import openpyxl
    wb = openpyxl.load_workbook(REPORT, read_only=True)
    ws = wb["岗位总表"]
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert hdr == HEADER_26, f"岗位总表表头与交付契约不一致\n实际: {hdr}"

def test_total_sheet_rows_and_links():
    import openpyxl
    wb = openpyxl.load_workbook(REPORT, read_only=True)
    ws = wb["岗位总表"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) >= 100, f"岗位总表行数异常: {len(rows)}"
    # 第 26 列 = 岗位链接，至少 50% 非空
    with_link = sum(1 for r in rows if r[25] not in (None, ""))
    assert with_link >= len(rows) // 2, f"岗位链接缺失过多: {with_link}/{len(rows)}"

def test_band_colors_present():
    import openpyxl
    wb = openpyxl.load_workbook(REPORT, read_only=True)
    ws = wb["岗位总表"]
    seen = set()
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        cell = row[0]
        fill = cell.fill
        if fill and fill.fgColor and fill.fgColor.rgb and fill.fgColor.rgb != "00000000":
            seen.add(str(fill.fgColor.rgb).replace("00", "", 1) if str(fill.fgColor.rgb).startswith("00") else str(fill.fgColor.rgb))
    assert seen & BAND_COLORS, f"未检测到任何档位底色: {seen}"

def test_tracker_exists():
    tp = os.path.join(ROOT, "data", "投递追踪.md")
    assert os.path.exists(tp), "投递追踪.md 缺失"

if __name__ == "__main__":
    test_report_exists()
    test_sheet_structure()
    test_total_sheet_header()
    test_total_sheet_rows_and_links()
    test_band_colors_present()
    test_tracker_exists()
    print("[PASS] delivery validation: report exists / 4 sheets / 26-col header / rows / links / band colors")